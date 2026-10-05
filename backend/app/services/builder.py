"""Source-preserving paper scaffolding, analysis integration, and exports."""
from __future__ import annotations

from html import escape
from io import BytesIO
import json
import re
from typing import Any

from pipeline.schemas import PaperRun
from pipeline.stages import StageState

from .store import LocalRunStore

LAYOUTS: dict[str, list[tuple[str, str]]] = {
    "imrad": [
        ("abstract", "Abstract"), ("introduction", "Introduction"),
        ("related_work", "Related Work"), ("methodology", "Methods"),
        ("experiments_results", "Results"), ("discussion", "Discussion"),
        ("conclusion", "Conclusion"), ("references", "References"),
    ],
    "conference": [
        ("abstract", "Abstract"), ("introduction", "Introduction"),
        ("related_work", "Related Work"), ("methodology", "Methodology"),
        ("experiments_results", "Experiments and Results"), ("discussion", "Discussion"),
        ("conclusion", "Conclusion"), ("references", "References"),
    ],
    "thesis": [
        ("abstract", "Abstract"), ("introduction", "Introduction"),
        ("related_work", "Literature Review"), ("methodology", "Methodology"),
        ("experiments_results", "Results"), ("discussion", "Discussion"),
        ("conclusion", "Conclusions"), ("references", "References"),
        ("appendices", "Appendices"),
    ],
    "generic": [
        ("abstract", "Abstract"), ("introduction", "Introduction"),
        ("related_work", "Related Work"), ("problem_statement", "Problem Statement"),
        ("objectives", "Objectives"), ("methodology", "Methodology"),
        ("dataset", "Dataset"), ("technologies", "Technologies"),
        ("experiments_results", "Experiments and Results"), ("discussion", "Discussion"),
        ("limitations", "Limitations"), ("future_work", "Future Work"),
        ("conclusion", "Conclusion"), ("references", "References"),
    ],
    "ieee": [
        ("abstract", "Abstract"), ("keywords", "Index Terms"),
        ("introduction", "Introduction"), ("problem_statement", "Problem Statement"),
        ("related_work", "Related Work"), ("objectives", "Objectives"),
        ("methodology", "Methodology"), ("dataset", "Dataset"),
        ("technologies", "Implementation Details"), ("experiments_results", "Experiments and Results"),
        ("discussion", "Discussion"), ("limitations", "Limitations"),
        ("future_work", "Future Work"), ("conclusion", "Conclusion"),
        ("references", "References"),
    ],
    "apa": [
        ("abstract", "Abstract"), ("keywords", "Keywords"),
        ("introduction", "Introduction"), ("problem_statement", "Problem Statement"),
        ("objectives", "Objectives"), ("related_work", "Literature Review"),
        ("methodology", "Method"), ("dataset", "Data and Materials"),
        ("technologies", "Technologies"), ("experiments_results", "Results"),
        ("discussion", "Discussion"), ("limitations", "Limitations"),
        ("future_work", "Future Directions"), ("conclusion", "Conclusion"),
        ("references", "References"),
    ],
    "mla": [
        ("abstract", "Abstract"), ("keywords", "Keywords"),
        ("introduction", "Introduction"), ("problem_statement", "Problem Statement"),
        ("related_work", "Related Work"), ("objectives", "Objectives"),
        ("methodology", "Methodology"), ("dataset", "Materials"),
        ("technologies", "Tools and Technologies"), ("experiments_results", "Results"),
        ("discussion", "Discussion"), ("limitations", "Limitations"),
        ("future_work", "Further Research"), ("conclusion", "Conclusion"),
        ("references", "Works Cited"),
    ],
    "chicago": [
        ("abstract", "Abstract"), ("keywords", "Keywords"),
        ("introduction", "Introduction"), ("problem_statement", "Problem Statement"),
        ("related_work", "Background and Related Work"), ("objectives", "Objectives"),
        ("methodology", "Methodology"), ("dataset", "Sources and Data"),
        ("technologies", "Tools and Technologies"), ("experiments_results", "Findings"),
        ("discussion", "Discussion"), ("limitations", "Limitations"),
        ("future_work", "Further Research"), ("conclusion", "Conclusion"),
        ("notes", "Notes"), ("references", "Bibliography"),
    ],
    "university": [
        ("institution", "Institution"), ("abstract", "Abstract"),
        ("objectives", "Objectives"), ("problem_statement", "Problem Statement"),
        ("introduction", "Introduction"), ("related_work", "Literature Review"),
        ("methodology", "Methodology"), ("dataset", "Dataset and Materials"),
        ("technologies", "Technologies"), ("experiments_results", "Results"),
        ("discussion", "Discussion"), ("limitations", "Limitations"),
        ("future_work", "Future Work"), ("conclusion", "Conclusion"),
        ("references", "References"),
    ],
}

INPUT_DEFAULTS: dict[str, Any] = {
    "title": "", "authors": [], "abstract": "", "institution": "", "keywords": [],
    "course": "", "instructor": "", "submission_date": "",
    "problem_statement": "", "objectives": "", "dataset": "", "technologies": "",
    "methodology": "",
    "experiments_results": "", "limitations": "", "future_work": "",
    "introduction": "", "related_work": "", "discussion": "", "conclusion": "",
    "sections": {}, "citations": [], "references": [],
}


def assemble_document(data: dict[str, Any]) -> dict[str, Any]:
    """Build an explicitly incomplete outline using only supplied text."""
    supplied = {**INPUT_DEFAULTS, **data}
    raw_sections = supplied.get("sections") or {}
    edits = supplied.get("generated_content") or {}
    normalized = {re.sub(r"[^a-z0-9]+", "_", key.lower()).strip("_"): text
                  for key, text in raw_sections.items()}
    normalized.update({
        "introduction": normalized.get("introduction", normalized.get("background", "")),
        "related_work": normalized.get("related_work", normalized.get("literature_review", "")),
        "methodology": normalized.get("methodology", normalized.get("methods", "")),
        "experiments_results": normalized.get("experiments_results", normalized.get("results", normalized.get("findings", ""))),
        "discussion": normalized.get("discussion", ""),
        "conclusion": normalized.get("conclusion", normalized.get("conclusions", "")),
        "references": normalized.get("references", ""),
        "appendices": normalized.get("appendices", normalized.get("appendix", "")),
    })

    aliases = {
        "abstract": "abstract",
        "introduction": "introduction",
        "related_work": "related_work",
        "problem_statement": "problem_statement",
        "objectives": "objectives",
        "methodology": "methodology",
        "dataset": "dataset",
        "technologies": "technologies",
        "experiments_results": "experiments_results",
        "discussion": "discussion",
        "limitations": "limitations",
        "future_work": "future_work",
        "conclusion": "conclusion",
        "institution": "institution",
    }
    for key, field in aliases.items():
        if supplied.get(field) and not normalized.get(key):
            normalized[key] = supplied[field]

    authors = supplied.get("authors") or []
    if isinstance(authors, str):
        authors = [author.strip() for author in authors.split(",") if author.strip()]
    keywords = supplied.get("keywords") or []
    if isinstance(keywords, str):
        keywords = [word.strip() for word in keywords.split(",") if word.strip()]
    references_input = supplied.get("references") or []
    if isinstance(references_input, str):
        references_input = [line for line in references_input.splitlines() if line.strip()]
    citations_input = supplied.get("citations") or []
    if isinstance(citations_input, str):
        citations_input = [line for line in citations_input.splitlines() if line.strip()]

    sections: list[dict[str, Any]] = []
    missing: list[str] = []
    if not supplied.get("title", "").strip():
        missing.append("title")
    if not authors:
        missing.append("authors")

    layout = LAYOUTS[supplied.get("format", "imrad")]
    layout_keys = {key for key, _ in layout}
    for key, heading in layout:
        if key in edits and isinstance(edits[key], str) and not edits[key].startswith("[Missing:"):
            text = edits[key]
        elif key == "references" and references_input:
            text = "\n".join(references_input)
        elif key == "keywords" and keywords:
            text = ", ".join(keywords)
        elif key == "abstract":
            text = supplied.get("abstract", "")
        else:
            text = normalized.get(key, "")

        if key == "references" and not text and not citations_input:
            text = "[Missing: references and citations. Add source details; none have been generated.]"
            missing.append("references and citations")
        elif key == "abstract" and not text:
            text = "[Missing: abstract]"
            missing.append("abstract")
        elif not text:
            text = f"[Missing: {heading}]"
            missing.append(heading.lower())
        if key == "references" and citations_input:
            citation_text = "Supplied citation markers (not verified or expanded):\n" + "\n".join(citations_input)
            if text.startswith("[Missing:"):
                text += "\n\n" + citation_text
            else:
                text += "\n\n" + citation_text
        sections.append({
            "key": key, "heading": heading, "content": text,
            "supplied": not text.startswith("[Missing:"),
        })

    extra_content = {
        **normalized,
        **{re.sub(r"[^a-z0-9]+", "_", key.lower()).strip("_"): value for key, value in raw_sections.items()},
        **{key: value for key, value in edits.items()},
    }
    aliases_to_canonical = {
        "background": "introduction", "literature_review": "related_work",
        "methods": "methodology", "results": "experiments_results",
        "findings": "experiments_results", "conclusions": "conclusion",
        "appendix": "appendices",
    }
    for key, value in extra_content.items():
        canonical_key = aliases_to_canonical.get(key, key)
        if canonical_key not in layout_keys and canonical_key != "institution" and isinstance(value, str) and value and not value.startswith("[Missing:"):
            sections.append({
                "key": key,
                "heading": key.replace("_", " ").title(),
                "content": value,
                "supplied": True,
            })

    return {
        "title": supplied.get("title", ""),
        "authors": authors,
        "institution": supplied.get("institution", ""),
        "keywords": keywords,
        "citations": citations_input,
        "generated_content": {section["key"]: section["content"] for section in sections},
        "sections": sections,
        "missing_items": missing,
        "notice": "Only user-supplied information is included. Missing items are marked; facts, results, citations, and references are not generated.",
    }


def generate_document(data: dict[str, Any]) -> dict[str, Any]:
    """Fill missing paper sections with clearly labeled, topic-grounded assumptions."""
    supplied = {**INPUT_DEFAULTS, **data}
    document = assemble_document({**supplied, "generated_content": {}})
    section_by_key = {section["key"]: section for section in document["sections"]}
    assumptions: list[str] = []

    def text_for(key: str) -> str:
        value = supplied.get(key, "")
        if isinstance(value, list):
            return "; ".join(str(item).strip() for item in value if str(item).strip())
        return str(value or "").strip()

    title = text_for("title")
    topic = title or text_for("problem_statement") or text_for("abstract") or "the research project"
    problem = text_for("problem_statement") or (
        f"the research problem described by the project topic, {topic}"
    )
    objectives = text_for("objectives") or (
        f"examine {topic}, describe a reproducible approach, and evaluate it using "
        "measurements collected from the actual project"
    )
    methodology = text_for("methodology") or (
        "A proposed workflow is to document the data source and preparation, define a "
        "baseline, train the selected method on a documented training partition, and "
        "evaluate it on held-out data using metrics appropriate to the task. The exact "
        "data split, algorithms, and parameters must be supplied from the implementation."
    )
    dataset = text_for("dataset") or (
        f"Assumed data requirement for {topic}: a dataset relevant to the stated research "
        "question will be needed. No source, record count, date range, location, variables, "
        "or license was supplied; these details must be replaced with the actual dataset."
    )
    technologies = text_for("technologies") or (
        "Assumed implementation option: use a documented programming environment and "
        "task-appropriate data-processing and modeling libraries. No specific software "
        "was reported as used in the supplied project details."
    )
    supplied_results = text_for("experiments_results")
    results = supplied_results or (
        "Assumption (hypothetical, not measured): the evaluation will compare the proposed "
        "approach with a clearly documented baseline using task-appropriate metrics. No "
        "scores, model ranking, sample counts, or experimental outcomes were supplied, so "
        "this draft does not report numerical results. Replace this statement with outputs "
        "from the completed experiment."
    )
    related_work = text_for("related_work") or (
        f"Assumption-based literature-review plan for {topic}: organize verified sources "
        "around the research problem, the methods used by prior work, and the evaluation "
        "criteria relevant to this project. These are review themes inferred from the "
        "supplied topic, not claims about specific publications. Add and verify real sources "
        "before treating this section as a literature review."
    )
    references = text_for("references")
    if not references:
        citations = text_for("citations")
        references = citations or (
            "Assumption placeholder (not a real citation): add a verified source relevant "
            f"to {topic}. Author, title, venue, year, pages, and DOI were not supplied and "
            "must not be treated as bibliographic facts."
        )
    citations = text_for("citations")

    abstract_result = (
        f"The supplied project details report: {supplied_results}"
        if supplied_results
        else "No experimental outcomes were supplied; results remain hypothetical and must be replaced with measured findings."
    )
    abstract_method = (
        methodology[:700].rstrip()
        + ("..." if len(methodology) > 700 else "")
    )
    inferred: dict[str, str] = {
        "abstract": (
            f"This study examines {topic} in relation to {problem}. Its objectives are to "
            f"{objectives}. The proposed methodology is {abstract_method} {abstract_result} "
            "This assumption-based abstract must be checked against the implemented project."
        ),
        "introduction": (
            f"This paper focuses on {topic}. The supplied project details identify the "
            f"following problem: {problem} The study is intended to {objectives}. "
            "This introduction is an assumption-based framing derived from the project "
            "description; add verified context and sources before submission."
        ),
        "related_work": related_work,
        "problem_statement": (
            f"Assumption-based framing: the central research problem is {problem}. This "
            "interpretation is inferred from the project title and supplied description "
            "and should be confirmed by the author."
        ),
        "objectives": (
            f"Assumed objectives to verify:\n1. Examine {topic}.\n"
            "2. Document the actual data and method used to address the problem.\n"
            "3. Evaluate the approach against a documented baseline using measured results.\n"
            "These objectives are inferred and may be edited to match the actual project."
        ),
        "methodology": f"Assumption-based proposed methodology: {methodology}",
        "dataset": f"Assumption-based dataset description: {dataset}",
        "technologies": f"Assumption-based implementation note: {technologies}",
        "experiments_results": results,
        "discussion": (
            f"Assumption-based discussion: interpret the completed evaluation in relation "
            f"to the objectives for {topic}. If the measured results support the approach, "
            "discuss the size and practical meaning of the observed change; if they do not, "
            "report that outcome and examine possible causes. No such outcome was supplied, "
            "so this draft makes no claim that either result occurred."
        ),
        "limitations": (
            "Assumption-based limitations to check: the data may not represent all relevant "
            "settings; measurement quality and missing values may affect outcomes; and "
            "performance on one evaluation split may not generalize. Confirm which of these "
            "limitations actually apply to the project and remove the rest."
        ),
        "future_work": (
            f"Assumption-based future work: validate the approach for {topic} on additional "
            "data or settings, compare it with further documented baselines, and report "
            "reproducible measurements. These are suggested next steps, not completed work."
        ),
        "conclusion": (
            f"Assumption-based conclusion: this project examines {topic} and is intended "
            f"to address {problem}. The stated objectives are to {objectives}. "
            "No effectiveness or deployment conclusion can be drawn until verified "
            "experimental results are supplied."
        ),
        "references": (
            "Supplied citation information (unverified; validate before use): " + references
            if text_for("citations")
            else references
        ),
        "notes": (
            "Assumption-based note entries (citation locations not supplied; map each "
            "note to its supporting passage before use):\n" + citations
            if citations
            else "No note text or citation locations were supplied. Add verified Chicago-style footnotes or endnotes and connect each note to the passage it supports."
        ),
    }

    for key, body in inferred.items():
        section = section_by_key.get(key)
        if section is None or not section["content"].startswith("[Missing:"):
            continue
        section["content"] = (
            "[ASSUMPTION-BASED DRAFT — verify before use]\n" + body
        )
        section["supplied"] = False
        assumptions.append(key)

    keywords = supplied.get("keywords") or []
    if isinstance(keywords, str):
        keywords = [word.strip() for word in keywords.split(",") if word.strip()]
    if not keywords and "keywords" in section_by_key:
        candidates = re.findall(r"[A-Za-z][A-Za-z-]{2,}", topic)
        stop_words = {"using", "with", "from", "into", "based", "study", "research", "and", "for"}
        inferred_keywords = list(dict.fromkeys(
            word.lower() for word in candidates if word.lower() not in stop_words
        ))[:6]
        if inferred_keywords:
            section_by_key["keywords"]["content"] = (
                "[ASSUMPTION-BASED DRAFT — verify before use]\n"
                + ", ".join(inferred_keywords)
            )
            assumptions.append("keywords")

    generated_content = {
        section["key"]: section["content"] for section in document["sections"]
    }
    document["generated_content"] = generated_content
    document["assumptions"] = assumptions
    document["notice"] = (
        "Generated prose is assumption-based and must be verified. Missing results and "
        "references are labeled as hypothetical or placeholders; no bibliographic details "
        "or measured outcomes are represented as facts."
    )
    return document


def paper_analysis(paper_id: str, store: LocalRunStore) -> PaperRun | None:
    """Read the actual run produced by the existing Enthesis analysis pipeline."""
    try:
        return store.get(paper_id)
    except (OSError, ValueError, TypeError):
        return None


def get_findings(run: PaperRun) -> list[dict[str, Any]]:
    """Normalize actual pipeline findings without synthesizing findings or evidence."""
    findings: list[dict[str, Any]] = []
    states = run.states or {}
    for module, result in (run.results or {}).items():
        if not isinstance(result, dict):
            continue
        state = states.get(module)
        if hasattr(state, "value"):
            state = state.value
        if state in {StageState.PENDING.value, StageState.RUNNING.value, StageState.FAILED.value}:
            continue
        module_findings = result.get("findings")
        if not isinstance(module_findings, list):
            continue
        module_evidence = result.get("evidence")
        confidence = result.get("confidence")
        if not isinstance(confidence, (int, float)) or isinstance(confidence, bool):
            confidence = None
        for index, item in enumerate(module_findings):
            if not isinstance(item, dict):
                continue
            evidence = item.get("evidence")
            if evidence is None and isinstance(module_evidence, list):
                evidence = module_evidence
            if evidence is None:
                evidence = []
            title = next((item.get(key) for key in ("title", "name", "label", "type")
                          if isinstance(item.get(key), str) and item.get(key).strip()),
                         "[No finding title supplied by analysis]")
            description = next((item.get(key) for key in ("description", "summary", "message", "text")
                                if isinstance(item.get(key), str) and item.get(key).strip()), None)
            if description is None:
                description = json.dumps(item, ensure_ascii=False, sort_keys=True)
            finding_confidence = item.get("confidence", confidence)
            if not isinstance(finding_confidence, (int, float)) or isinstance(finding_confidence, bool):
                finding_confidence = None
            affected_section = next((item.get(key) for key in ("affected_section", "section")
                                     if isinstance(item.get(key), str) and item.get(key).strip()), None)
            recommendation = next((item.get(key) for key in ("recommended_action", "suggested_action", "suggestion")
                                   if isinstance(item.get(key), str) and item.get(key).strip()),
                                  "Review this analysis finding against the manuscript; the analysis supplied no specific recommended action.")
            why_detected = next((item.get(key) for key in ("why_detected", "detection_reason")
                                 if isinstance(item.get(key), str) and item.get(key).strip()),
                                f"Reported by the existing {module} analysis; no further detection rationale was supplied.")
            findings.append({
                "id": f"{module}:{index}",
                "title": title,
                "description": description,
                "why_detected": why_detected,
                "evidence": evidence,
                "confidence": finding_confidence,
                "affected_section": affected_section,
                "recommended_action": recommendation,
                "source_module": module,
            })
    return findings


def quality_dashboard(run: PaperRun) -> dict[str, Any]:
    """Summarize real pipeline outputs; omit synthetic scores and empty findings."""
    modules = []
    for module, result in (run.results or {}).items():
        if not isinstance(result, dict):
            continue
        state = run.states.get(module, "unknown")
        state = state.value if hasattr(state, "value") else state
        actual_findings = result.get("findings")
        if not isinstance(actual_findings, list):
            actual_findings = []
        entry: dict[str, Any] = {
            "module": module,
            "status": state,
            "findings_count": len(actual_findings),
        }
        if "metrics" in result:
            entry["metrics"] = result["metrics"]
        if "confidence" in result:
            entry["confidence"] = result["confidence"]
        modules.append(entry)
    return {
        "paper_id": run.paper_id,
        "modules": modules,
        "findings_count": len(get_findings(run)),
        "notice": "Dashboard values are reported by the existing analysis modules; no aggregate quality score is inferred.",
    }


def export_document(document: dict[str, Any], file_format: str) -> tuple[bytes, str]:
    """Render persisted builder content into a valid portable document."""
    sections = document.get("sections", [])
    title = document.get("title") or "[Missing: title]"
    section_keys = {section.get("key") for section in sections}
    if file_format == "markdown":
        blocks = [f"# {title}"]
        authors = document.get("authors") or []
        if authors:
            blocks.extend(["", ", ".join(authors)])
        elif "authors" in document.get("missing_items", []):
            blocks.extend(["", "[Missing: authors]"])
        if document.get("institution") and "institution" not in section_keys:
            blocks.extend(["", document["institution"]])
        if document.get("keywords") and "keywords" not in section_keys:
            blocks.extend(["", "**Keywords:** " + ", ".join(document["keywords"])])
        for section in sections:
            blocks.extend(["", f"## {section['heading']}", "", section["content"]])
        if document.get("keywords"):
            blocks.extend(["", "**Keywords:** " + ", ".join(document["keywords"])])
        return ("\n".join(blocks) + "\n").encode("utf-8"), "text/markdown; charset=utf-8"

    if file_format == "latex":
        def tex(value: Any) -> str:
            value = str(value)
            replacements = {
                "\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$",
                "#": r"\#", "_": r"\_", "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}",
                "^": r"\textasciicircum{}",
            }
            return "".join(replacements.get(char, char) for char in value)

        blocks = [r"\documentclass{article}", r"\usepackage[utf8]{inputenc}", r"\begin{document}",
                  rf"\title{{{tex(title)}}}"]
        if document.get("authors"):
            blocks.append(r"\author{" + tex(", ".join(document["authors"])) + "}")
        blocks.extend([r"\maketitle"])
        if document.get("institution") and "institution" not in section_keys:
            blocks.append(tex(document["institution"]))
        if document.get("keywords") and "keywords" not in section_keys:
            blocks.extend([r"\textbf{Keywords:} " + tex(", ".join(document["keywords"]))])
        if "authors" in document.get("missing_items", []):
            blocks.append(r"\textit{[Missing: authors]}")
        for section in sections:
            blocks.extend([rf"\section{{{tex(section['heading'])}}}", tex(section["content"])])
        blocks.append(r"\end{document}")
        return ("\n".join(blocks) + "\n").encode("utf-8"), "application/x-tex; charset=utf-8"

    if file_format == "docx":
        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        doc = Document()
        input_details = document.get("input", {})
        paper_format = input_details.get("format", "generic")
        is_ieee = paper_format == "ieee"
        is_apa = paper_format == "apa"
        is_mla = paper_format == "mla"
        is_chicago = paper_format == "chicago"
        is_author_paper = is_apa or is_mla or is_chicago
        paper_section = doc.sections[0]
        paper_section.page_width = Inches(8.5 if is_ieee or is_author_paper else 8.27)
        paper_section.page_height = Inches(11 if is_ieee or is_author_paper else 11.69)
        margin = Inches(0.7 if is_ieee else 1)
        paper_section.top_margin = margin
        paper_section.bottom_margin = margin
        paper_section.left_margin = margin
        paper_section.right_margin = margin

        normal_style = doc.styles["Normal"]
        normal_style.font.name = "Times New Roman"
        normal_style.font.size = Pt(10 if is_ieee else 12)
        normal_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        normal_style.paragraph_format.line_spacing = 1.0 if is_ieee else (2.0 if is_author_paper else 1.5)
        normal_style.paragraph_format.space_after = Pt(6)
        normal_style.paragraph_format.first_line_indent = Inches(0 if is_ieee else 0.5)

        if is_apa or is_chicago:
            header = paper_section.header.paragraphs[0]
            header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            page_field = OxmlElement("w:fldSimple")
            page_field.set(qn("w:instr"), "PAGE")
            header._p.append(page_field)
        elif is_mla:
            header = paper_section.header.paragraphs[0]
            header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            header.add_run(", ".join(document.get("authors") or []) or "[Missing: author]")
            page_field = OxmlElement("w:fldSimple")
            page_field.set(qn("w:instr"), "PAGE")
            header._p.append(page_field)

        authors = ", ".join(document.get("authors") or []) or "[Missing: authors]"
        if is_mla:
            for value in (
                authors,
                input_details.get("instructor") or "[Missing: instructor]",
                input_details.get("course") or "[Missing: course]",
                input_details.get("submission_date") or "[Missing: submission date]",
            ):
                doc.add_paragraph(value)
            title_paragraph = doc.add_paragraph()
            title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title_paragraph.add_run(title)
        else:
            if is_apa or is_chicago:
                for _ in range(3 if is_apa else 5):
                    doc.add_paragraph()
            title_paragraph = doc.add_paragraph()
            title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title_run = title_paragraph.add_run(title)
            title_run.bold = True
            title_run.font.name = "Times New Roman"
            title_run.font.size = Pt(22 if is_ieee else 18)
            author_paragraph = doc.add_paragraph(authors)
            author_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if document.get("institution") and "institution" not in section_keys:
                institution_paragraph = doc.add_paragraph(document["institution"])
                institution_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if is_apa:
                for field, missing_label in (
                    ("course", "course"),
                    ("instructor", "instructor"),
                    ("submission_date", "due date"),
                ):
                    value = input_details.get(field) or f"[Missing: {missing_label}]"
                    metadata_paragraph = doc.add_paragraph(value)
                    metadata_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if is_chicago and input_details.get("submission_date"):
                date_paragraph = doc.add_paragraph(input_details["submission_date"])
                date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if is_apa or is_chicago:
            doc.add_page_break()

        if is_ieee:
            two_column_section = doc.add_section(WD_SECTION.CONTINUOUS)
            columns = two_column_section._sectPr.xpath("./w:cols")
            if not columns:
                columns_element = OxmlElement("w:cols")
                two_column_section._sectPr.append(columns_element)
                columns = [columns_element]
            columns[0].set(qn("w:num"), "2")
            columns[0].set(qn("w:space"), "360")

        ieee_body_index = 0
        for section in sections:
            key = section["key"]
            heading = section["heading"]
            if key == "keywords":
                heading = "Index Terms" if is_ieee else "Keywords"
            if key == "references":
                heading = {
                    "mla": "Works Cited",
                    "chicago": "Bibliography",
                }.get(paper_format, heading)
            heading_paragraph = doc.add_paragraph()
            heading_paragraph.paragraph_format.keep_with_next = True
            heading_paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if is_ieee or is_apa or key == "abstract" else WD_ALIGN_PARAGRAPH.LEFT
            )
            heading_run = heading_paragraph.add_run(heading.upper() if is_ieee else heading)
            heading_run.bold = True
            heading_run.font.name = "Times New Roman"
            heading_run.font.size = Pt(10 if is_ieee else 12)

            if is_ieee and key not in {"abstract", "keywords", "references"}:
                ieee_body_index += 1
                roman = (
                    "I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X",
                    "XI", "XII", "XIII", "XIV", "XV",
                )
                heading_run.text = f"{roman[ieee_body_index - 1]}. {heading.upper()}" if ieee_body_index <= len(roman) else f"{ieee_body_index}. {heading.upper()}"

            if key == "references" and not section["content"].startswith("[Missing:"):
                references = [line.strip() for line in section["content"].splitlines() if line.strip()]
                for index, reference in enumerate(references, 1):
                    paragraph = doc.add_paragraph()
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    paragraph.paragraph_format.left_indent = Inches(0.2)
                    paragraph.paragraph_format.first_line_indent = Inches(-0.2) if not is_ieee else Inches(0)
                    paragraph.paragraph_format.space_after = Pt(3)
                    prefix = f"[{index}] " if is_ieee and not re.match(r"^\[\d+\]", reference) else ""
                    paragraph.add_run(prefix + reference)
                continue

            paragraphs = section["content"].splitlines() or [""]
            for text in paragraphs:
                paragraph = doc.add_paragraph(text)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                if key == "abstract":
                    paragraph.paragraph_format.first_line_indent = Inches(0)

        output = BytesIO()
        doc.save(output)
        return output.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    if file_format == "pdf":
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

        output = BytesIO()
        pdf = SimpleDocTemplate(output, pagesize=A4, rightMargin=inch, leftMargin=inch,
                                topMargin=inch, bottomMargin=inch)
        styles = getSampleStyleSheet()
        flow = [Paragraph(escape(title), styles["Title"]), Spacer(1, 12)]
        if document.get("authors"):
            flow.extend([Paragraph(escape(", ".join(document["authors"])), styles["Normal"]), Spacer(1, 12)])
        elif "authors" in document.get("missing_items", []):
            flow.extend([Paragraph("[Missing: authors]", styles["Normal"]), Spacer(1, 12)])
        if document.get("institution") and "institution" not in section_keys:
            flow.extend([Paragraph(escape(document["institution"]), styles["Normal"]), Spacer(1, 12)])
        if document.get("keywords") and "keywords" not in section_keys:
            flow.extend([Paragraph("<b>Keywords:</b> " + escape(", ".join(document["keywords"])), styles["Normal"]), Spacer(1, 12)])
        for section in sections:
            flow.append(Paragraph(escape(section["heading"]), styles["Heading1"]))
            text = escape(section["content"]).replace("\n", "<br/>")
            flow.extend([Paragraph(text, styles["BodyText"]), Spacer(1, 10)])
        pdf.build(flow)
        return output.getvalue(), "application/pdf"

    raise ValueError(f"Unsupported export format: {file_format}")
