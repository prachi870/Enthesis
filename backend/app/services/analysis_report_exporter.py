"""Render persisted analysis reports as downloadable documents."""
import json
from html import escape
from io import BytesIO
from typing import Any


def _display(value: Any) -> str:
    if value is None or value == "":
        return "Not supplied"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, indent=2)
    return str(value)


def _report_markdown(report: dict) -> str:
    summary = report.get("summary", {})
    lines = [
        f"# Research Analysis Report: {report.get('paper_title') or 'Untitled paper'}",
        "",
        f"- Paper ID: {report.get('paper_id', 'Not supplied')}",
        f"- Version: {report.get('version', 'Not supplied')}",
        f"- Generated: {report.get('generated_at', 'Not supplied')}",
        f"- Analysis status: {report.get('analysis_status', 'Not supplied')}",
        "",
        "## Summary",
        "",
        f"- Total findings: {summary.get('total_findings', 0)}",
        f"- Requiring investigation: {summary.get('requiring_investigation', 0)}",
        f"- Evidence passages: {summary.get('evidence_count', 0)}",
        "",
        "## Module status",
        "",
    ]
    for key, module in report.get("modules", {}).items():
        lines.append(f"- **{key.replace('_', ' ').title()}**: {module.get('status', 'Not supplied')}")
    lines.extend(["", "## Findings", ""])
    findings = report.get("findings", [])
    if not findings:
        lines.append("No findings were returned by the available analysis modules.")
    for index, finding in enumerate(findings, 1):
        lines.extend([
            f"### {index}. {finding.get('title') or 'Untitled finding'}",
            "",
            f"- Module: {_display(finding.get('module'))}",
            f"- Status: {_display(finding.get('status'))}",
            f"- Affected section: {_display(finding.get('affected_section'))}",
            f"- Confidence: {_display(finding.get('confidence'))}",
            "",
            f"**Description**  \n{_display(finding.get('description'))}",
            "",
            f"**Detailed explanation**  \n{_display(finding.get('detailed_explanation'))}",
            "",
            f"**Why detected**  \n{_display(finding.get('why_detected'))}",
            "",
            f"**Evidence**  \n{_display(finding.get('evidence_text'))}",
            "",
            f"**Related research**  \n{_display(finding.get('related_research'))}",
            "",
            f"**Recommended action**  \n{_display(finding.get('recommended_action'))}",
            "",
            f"**Analysis limitation**  \n{_display(finding.get('analysis_limitation'))}",
            "",
        ])
    lines.extend(["## Limitations", ""])
    lines.extend(f"- {item}" for item in report.get("limitations", []))
    lines.append("")
    return "\n".join(lines)


def export_analysis_report(report: dict, file_format: str) -> tuple[bytes, str, str]:
    """Return document bytes, media type, and file extension for a report."""
    if file_format == "markdown":
        return _report_markdown(report).encode("utf-8"), "text/markdown; charset=utf-8", "md"
    if file_format == "json":
        return (
            json.dumps(report, ensure_ascii=False, indent=2).encode("utf-8"),
            "application/json",
            "json",
        )

    if file_format == "docx":
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.shared import Inches, Pt
        except ImportError as exc:
            raise ImportError("python-docx is required for DOCX export.") from exc

        document = Document()
        for section in document.sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.85)
            section.right_margin = Inches(0.85)
        title = document.add_heading(
            f"Research Analysis Report: {report.get('paper_title') or 'Untitled paper'}",
            level=0,
        )
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph(
            f"Version {report.get('version', 'Not supplied')} · "
            f"Generated {report.get('generated_at', 'Not supplied')}"
        )
        document.add_paragraph(f"Analysis status: {report.get('analysis_status', 'Not supplied')}")
        document.add_heading("Summary", level=1)
        summary = report.get("summary", {})
        for label, key in (
            ("Total findings", "total_findings"),
            ("Requiring investigation", "requiring_investigation"),
            ("Evidence passages", "evidence_count"),
        ):
            document.add_paragraph(f"{label}: {summary.get(key, 0)}")
        document.add_heading("Module status", level=1)
        for key, module in report.get("modules", {}).items():
            document.add_paragraph(
                f"{key.replace('_', ' ').title()}: {module.get('status', 'Not supplied')}",
                style="List Bullet",
            )
        document.add_heading("Findings", level=1)
        findings = report.get("findings", [])
        if not findings:
            document.add_paragraph("No findings were returned by the available analysis modules.")
        for index, finding in enumerate(findings, 1):
            document.add_heading(
                f"{index}. {finding.get('title') or 'Untitled finding'}",
                level=2,
            )
            fields = (
                ("Module", finding.get("module")),
                ("Status", finding.get("status")),
                ("Affected section", finding.get("affected_section")),
                ("Confidence", finding.get("confidence")),
                ("Description", finding.get("description")),
                ("Detailed explanation", finding.get("detailed_explanation")),
                ("Why detected", finding.get("why_detected")),
                ("Evidence", finding.get("evidence_text")),
                ("Related research", finding.get("related_research")),
                ("Recommended action", finding.get("recommended_action")),
                ("Analysis limitation", finding.get("analysis_limitation")),
            )
            for label, value in fields:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.add_run(f"{label}: ").bold = True
                paragraph.add_run(_display(value))
        document.add_heading("Limitations", level=1)
        for limitation in report.get("limitations", []):
            document.add_paragraph(limitation, style="List Bullet")
        output = BytesIO()
        document.save(output)
        return (
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "docx",
        )

    if file_format == "pdf":
        try:
            from reportlab.lib import colors
            from reportlab.lib.enums import TA_CENTER
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import inch
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        except ImportError as exc:
            raise ImportError("reportlab is required for PDF export.") from exc

        output = BytesIO()
        document = SimpleDocTemplate(
            output, pagesize=A4, rightMargin=0.75 * inch, leftMargin=0.75 * inch,
            topMargin=0.7 * inch, bottomMargin=0.7 * inch,
        )
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "AnalysisReportTitle", parent=styles["Title"], alignment=TA_CENTER,
            textColor=colors.HexColor("#283593"), spaceAfter=14,
        )
        body_style = ParagraphStyle(
            "AnalysisReportBody", parent=styles["BodyText"], leading=14, spaceAfter=6,
        )
        story = [
            Paragraph(escape(f"Research Analysis Report: {report.get('paper_title') or 'Untitled paper'}"), title_style),
            Paragraph(escape(f"Version {report.get('version', 'Not supplied')} · Generated {report.get('generated_at', 'Not supplied')}"), body_style),
            Paragraph(escape(f"Analysis status: {report.get('analysis_status', 'Not supplied')}"), body_style),
            Spacer(1, 8),
            Paragraph("Summary", styles["Heading1"]),
        ]
        summary = report.get("summary", {})
        for label, key in (
            ("Total findings", "total_findings"),
            ("Requiring investigation", "requiring_investigation"),
            ("Evidence passages", "evidence_count"),
        ):
            story.append(Paragraph(escape(f"{label}: {summary.get(key, 0)}"), body_style))
        story.append(Paragraph("Module status", styles["Heading1"]))
        for key, module in report.get("modules", {}).items():
            story.append(Paragraph(
                escape(f"{key.replace('_', ' ').title()}: {module.get('status', 'Not supplied')}"),
                body_style,
            ))
        story.append(Paragraph("Findings", styles["Heading1"]))
        findings = report.get("findings", [])
        if not findings:
            story.append(Paragraph("No findings were returned by the available analysis modules.", body_style))
        for index, finding in enumerate(findings, 1):
            story.append(Paragraph(
                escape(f"{index}. {finding.get('title') or 'Untitled finding'}"),
                styles["Heading2"],
            ))
            for label, field in (
                ("Module", "module"),
                ("Status", "status"),
                ("Affected section", "affected_section"),
                ("Confidence", "confidence"),
                ("Description", "description"),
                ("Detailed explanation", "detailed_explanation"),
                ("Why detected", "why_detected"),
                ("Evidence", "evidence_text"),
                ("Related research", "related_research"),
                ("Recommended action", "recommended_action"),
                ("Analysis limitation", "analysis_limitation"),
            ):
                text = escape(_display(finding.get(field))).replace("\n", "<br/>")
                story.append(Paragraph(f"<b>{escape(label)}:</b> {text}", body_style))
        story.append(Paragraph("Limitations", styles["Heading1"]))
        for limitation in report.get("limitations", []):
            story.append(Paragraph(f"• {escape(str(limitation))}", body_style))
        document.build(story)
        return output.getvalue(), "application/pdf", "pdf"

    raise ValueError(f"Unsupported analysis report format: {file_format}")
