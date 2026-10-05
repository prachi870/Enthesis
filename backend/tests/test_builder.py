from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pypdf import PdfReader
from sqlalchemy import create_engine, inspect
from sqlalchemy.orm import sessionmaker
import pytest
import time

from backend.app.api import builder as builder_api
from backend.app.api.builder import router
from backend.app.database import Base, get_db
from backend.app.services.store import LocalRunStore
from pipeline.schemas import PaperRun
from pipeline.stages import StageState


def make_client(tmp_path, monkeypatch):
    db_path = tmp_path / "builder.sqlite"
    engine = create_engine(f"sqlite:///{db_path}")
    Base.metadata.create_all(bind=engine)
    session_factory = sessionmaker(bind=engine, autocommit=False, autoflush=False)
    monkeypatch.setattr(builder_api, "SessionLocal", session_factory)

    def db_dependency():
        db = session_factory()
        try:
            yield db
        finally:
            db.close()

    app = FastAPI()
    app.include_router(router, prefix="/api/v1/builder")
    app.include_router(router, prefix="/api/v1/research-papers")
    app.dependency_overrides[get_db] = db_dependency
    paper_store = LocalRunStore(str(tmp_path / "paper_runs"))
    monkeypatch.setattr(builder_api, "analysis_store", paper_store)
    return TestClient(app), engine, paper_store


def test_draft_persistence_and_immutable_versions(tmp_path, monkeypatch):
    client, engine, _ = make_client(tmp_path, monkeypatch)
    try:
        response = client.post("/api/v1/builder/drafts", json={
            "title": "A supplied study of coastal sensor calibration",
            "authors": ["Mina Shah"],
            "abstract": "We evaluate calibration drift using observations collected at three coastal stations.",
            "sections": {
                "introduction": "The study examines sensor drift.",
                "methods": "Sensors were calibrated monthly.",
                "results": "The supplied observations show a median drift of 0.8 units.",
            },
            "references": ["Shah, M. (2025). Sensor calibration notes."],
            "format": "imrad",
        })
        assert response.status_code == 201, response.text
        draft = response.json()
        draft_id = draft["draft_id"]
        assert draft["version"] == 1
        assert draft["document"]["sections"][0]["content"] == (
            "We evaluate calibration drift using observations collected at three coastal stations."
        )
        assert "experiments_results" in [section["key"] for section in draft["document"]["sections"]]
        assert "discussion" in draft["document"]["missing_items"]

        update = client.put(f"/api/v1/builder/drafts/{draft_id}", json={
            "sections": {"discussion": "The observed drift is discussed here, without adding new data."},
            "format": "thesis",
        })
        assert update.status_code == 200, update.text
        assert update.json()["version"] == 2
        assert update.json()["format"] == "thesis"
        assert [section["heading"] for section in update.json()["document"]["sections"]][-2:] == [
            "References", "Appendices"
        ]
        persisted = client.get(f"/api/v1/builder/drafts/{draft_id}")
        assert persisted.json()["version"] == 2
        assert persisted.json()["document"]["sections"][-2]["content"].startswith(
            "Shah, M. (2025). Sensor calibration notes."
        )

        versions = client.get(f"/api/v1/builder/drafts/{draft_id}/versions")
        assert versions.status_code == 200
        assert [item["version"] for item in versions.json()["versions"]] == [1, 2]
        old = client.get(f"/api/v1/builder/drafts/{draft_id}/versions/1")
        assert old.json()["document"]["sections"][-1]["heading"] == "References"
        assert old.json()["format"] == "imrad"
    finally:
        client.close()
        engine.dispose()


def test_research_papers_client_contract_persists_supplied_fields(tmp_path, monkeypatch):
    client, engine, _ = make_client(tmp_path, monkeypatch)
    try:
        created = client.post("/api/v1/research-papers", json={
            "title": "Acoustic monitoring in wetland habitats",
            "authors": "Mina Shah",
            "institution": "North Shore Institute",
            "abstract": "The researcher supplied this summary.",
            "keywords": "wetland, acoustic monitoring",
            "introduction": "The supplied introduction describes seasonal sound changes.",
            "related_work": "The supplied review cites earlier monitoring work.",
            "problem_statement": "The supplied question concerns detection of seasonal sound changes.",
            "objectives": "Compare two recording periods.",
            "methodology": "Recordings were grouped by month.",
            "dataset": "A researcher-provided set of 40 recordings.",
            "technologies": "Field microphones",
            "experiments_results": "The supplied result is a 12 percent change.",
            "discussion": "The researcher supplied discussion text.",
            "limitations": "Only the supplied sample set is included.",
            "future_work": "The researcher plans to add locations.",
            "conclusion": "The supplied conclusion summarizes observations.",
            "references": "Shah (2025), Field recording methods.",
            "format": "generic",
        })
        assert created.status_code == 201, created.text
        draft = created.json()
        draft_id = draft["paper_id"]
        assert draft["id"] == draft_id
        assert draft["details"]["institution"] == "North Shore Institute"
        assert "12 percent change" in draft["generated_content"]["experiments_results"]
        assert created.json()["generated_content"]["discussion"] == "The researcher supplied discussion text."

        fetched = client.get(f"/api/v1/research-papers/{draft_id}")
        assert fetched.status_code == 200
        assert fetched.json()["details"]["authors"] == "Mina Shah"
        assert fetched.json()["generated_content"]["problem_statement"].startswith(
            "The supplied question"
        )
        assert client.get("/api/v1/research-papers").status_code == 200

        updated = client.put(f"/api/v1/research-papers/{draft_id}", json={
            **draft["details"],
            "format": "ieee",
            "generated_content": {**draft["generated_content"], "discussion": "Only user-authored text."},
        })
        assert updated.status_code == 200, updated.text
        assert updated.json()["version"] == 2
        assert updated.json()["format"] == "ieee"
        assert updated.json()["generated_content"]["discussion"] == "Only user-authored text."
        versions = client.get(f"/api/v1/research-papers/{draft_id}/versions")
        assert [item["version"] for item in versions.json()["versions"]] == [1, 2]

        regenerated = client.post(f"/api/v1/research-papers/{draft_id}/sections/discussion/regenerate")
        assert regenerated.json()["section_content"] == "Only user-authored text."
        assert "No new claims" in regenerated.json()["notice"]
        exported = client.get(f"/api/v1/research-papers/{draft_id}/export?format=markdown")
        assert exported.status_code == 200
        status = client.post(f"/api/v1/research-papers/{draft_id}/export/status", json={"status": "downloaded"})
        assert status.json()["export_status"]["downloaded"] is True
    finally:
        client.close()
        engine.dispose()


def test_research_paper_generate_endpoint_returns_generated_sections(tmp_path, monkeypatch):
    client, engine, _ = make_client(tmp_path, monkeypatch)
    try:
        created = client.post("/api/v1/research-papers", json={
            "title": "Coastal sensor calibration",
            "authors": "Mina Shah",
            "abstract": "This study evaluates sensor drift using observations supplied by the researcher.",
            "problem_statement": "Sensor readings may drift over time.",
            "methodology": "Sensors were calibrated monthly.",
            "experiments_results": "The supplied observations show a median drift of 0.8 units.",
            "format": "generic",
        })
        assert created.status_code == 201, created.text
        draft_id = created.json()["paper_id"]
        assert created.json()["is_generated"] is False

        saved_with_stale_section = client.put(f"/api/v1/research-papers/{draft_id}", json={
            **created.json()["details"],
            "format": "generic",
            "generated_content": {"abstract": "Stale section content from an earlier save."},
        })
        assert saved_with_stale_section.status_code == 200, saved_with_stale_section.text

        generated = client.post(f"/api/v1/research-papers/{draft_id}/generate")
        assert generated.status_code == 200, generated.text
        assert generated.json()["is_generated"] is True
        document = generated.json()["generated_content"]
        assert document["abstract"] == (
            "This study evaluates sensor drift using observations supplied by the researcher."
        )
        assert document["experiments_results"] == (
            "The supplied observations show a median drift of 0.8 units."
        )
        assert document["introduction"].startswith("[ASSUMPTION-BASED DRAFT")
        assert document["references"].startswith("[ASSUMPTION-BASED DRAFT")
        assert "Assumption placeholder (not a real citation)" in document["references"]
        assert "related_work" in generated.json()["assumptions"]
        assert generated.json()["version"] == 3
    finally:
        client.close()
        engine.dispose()


def test_exports_are_valid_pdf_docx_markdown_and_latex(tmp_path, monkeypatch):
    client, engine, _ = make_client(tmp_path, monkeypatch)
    try:
        created = client.post("/api/v1/builder/drafts", json={
            "title": "Field observations & sensor drift",
            "authors": ["Mina Shah"],
            "abstract": "Summary supplied by the researcher.",
            "sections": {
                "introduction": "The user provided this introduction.",
                "methods": "The user provided this method.",
                "results": "The user provided these results.",
            },
            "format": "conference",
        }).json()
        draft_id = created["draft_id"]

        pdf = client.get(f"/api/v1/builder/drafts/{draft_id}/export?format=pdf")
        assert pdf.status_code == 200
        parsed_pdf = PdfReader(BytesIO(pdf.content))
        assert len(parsed_pdf.pages) >= 1
        assert "Field observations" in "".join(page.extract_text() or "" for page in parsed_pdf.pages)

        docx = client.get(f"/api/v1/builder/drafts/{draft_id}/export?format=docx")
        parsed_docx = Document(BytesIO(docx.content))
        docx_text = "\n".join(paragraph.text for paragraph in parsed_docx.paragraphs)
        assert "Field observations & sensor drift" in docx_text
        assert "Experiments and Results" in docx_text
        assert "[Missing: Discussion]" in docx_text
        body_paragraph = next(p for p in parsed_docx.paragraphs if "The user provided this introduction." in p.text)
        assert body_paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY

        markdown = client.get(f"/api/v1/builder/drafts/{draft_id}/export?format=markdown")
        assert markdown.headers["content-type"].startswith("text/markdown")
        assert "# Field observations & sensor drift" in markdown.text
        assert "[Missing: references and citations" in markdown.text

        latex = client.get(f"/api/v1/builder/drafts/{draft_id}/export?format=latex")
        assert latex.headers["content-disposition"].endswith(".tex\"")
        assert r"\documentclass{article}" in latex.text
        assert r"Field observations \& sensor drift" in latex.text
        assert r"\begin{document}" in latex.text and r"\end{document}" in latex.text
        assert client.get(f"/api/v1/builder/drafts/{draft_id}/export?format=html").status_code == 422
    finally:
        client.close()
        engine.dispose()


def test_ieee_docx_export_uses_two_columns_and_keeps_supplied_text(tmp_path, monkeypatch):
    client, engine, _ = make_client(tmp_path, monkeypatch)
    try:
        created = client.post("/api/v1/research-papers", json={
            "title": "Researcher-supplied title",
            "authors": "Mina Shah",
            "abstract": "Researcher-supplied abstract.",
            "introduction": "Researcher-supplied introduction text.",
            "methodology": "Researcher-supplied method text.",
            "references": "Shah (2025). Researcher-supplied reference.",
            "format": "ieee",
        })
        assert created.status_code == 201, created.text
        generated = client.post(f"/api/v1/research-papers/{created.json()['paper_id']}/generate")
        assert generated.status_code == 200

        exported = client.get(
            f"/api/v1/research-papers/{created.json()['paper_id']}/export?format=docx"
        )
        assert exported.status_code == 200
        document = Document(BytesIO(exported.content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        assert "I. INTRODUCTION" in text
        assert "Researcher-supplied introduction text." in text
        assert "[1] Shah (2025). Researcher-supplied reference." in text

        columns = document.sections[-1]._sectPr.xpath("./w:cols")[0]
        assert columns.get(qn("w:num")) == "2"
        body_paragraph = next(p for p in document.paragraphs if "Researcher-supplied introduction text." in p.text)
        assert body_paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
        assert body_paragraph.style.font.name == "Times New Roman"
    finally:
        client.close()
        engine.dispose()


@pytest.mark.parametrize(
    ("paper_format", "expected_heading"),
    [("apa", "References"), ("mla", "Works Cited"), ("chicago", "Bibliography")],
)
def test_style_specific_docx_front_matter_and_spacing(
    tmp_path, monkeypatch, paper_format, expected_heading
):
    client, engine, _ = make_client(tmp_path, monkeypatch)
    try:
        created = client.post("/api/v1/research-papers", json={
            "title": "A supplied style test",
            "authors": "Mina Shah",
            "institution": "North Shore Institute",
            "course": "Research Methods",
            "instructor": "Dr. Lee",
            "submission_date": "September 30, 2026",
            "abstract": "A supplied abstract.",
            "references": "Shah, M. Supplied reference details.",
            "format": paper_format,
        })
        assert created.status_code == 201, created.text
        generated = client.post(
            f"/api/v1/research-papers/{created.json()['paper_id']}/generate"
        )
        assert generated.status_code == 200, generated.text

        exported = client.get(
            f"/api/v1/research-papers/{created.json()['paper_id']}/export?format=docx"
        )
        assert exported.status_code == 200, exported.text
        document = Document(BytesIO(exported.content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        assert expected_heading in text
        assert document.sections[0].top_margin.inches == pytest.approx(1)
        assert document.styles["Normal"].paragraph_format.line_spacing == 2.0
        if paper_format == "apa":
            assert "Dr. Lee" in text
            assert "Research Methods" in text
            assert any("Abstract" == paragraph.text for paragraph in document.paragraphs)
            assert any(
                break_element.get(qn("w:type")) == "page"
                for paragraph in document.paragraphs
                for break_element in paragraph._p.xpath(".//w:br")
            )
        if paper_format == "mla":
            assert "Dr. Lee" in text
            assert "Research Methods" in text
            assert "[Missing: course]" not in text
            assert document.sections[0].header.paragraphs[0].text.startswith("Mina Shah")
    finally:
        client.close()
        engine.dispose()


def test_generate_completes_missing_sections_with_labeled_assumptions(tmp_path, monkeypatch):
    client, engine, _ = make_client(tmp_path, monkeypatch)
    try:
        created = client.post("/api/v1/research-papers", json={
            "problem_statement": "Predict waste generation for better collection planning.",
            "format": "generic",
        })
        assert created.status_code == 201, created.text
        draft_id = created.json()["paper_id"]
        generated = client.post(f"/api/v1/research-papers/{draft_id}/generate")
        assert generated.status_code == 200, generated.text

        result = generated.json()
        sections = result["generated_content"]
        assert result["title"].startswith("Assumption-based study of")
        assert result["details"]["authors"] == ""
        assert result["assumptions"]
        assert all("[ASSUMPTION-BASED DRAFT" in text for key, text in sections.items()
                   if key in result["assumptions"] and key != "title")
        assert "hypothetical, not measured" in sections["experiments_results"]
        assert "not a real citation" in sections["references"]
        assert all(key in sections for key in (
            "abstract", "introduction", "related_work", "methodology",
            "experiments_results", "discussion", "conclusion", "references",
        ))
    finally:
        client.close()
        engine.dispose()


def test_builder_uses_existing_analysis_findings_and_metrics(tmp_path, monkeypatch):
    client, engine, paper_store = make_client(tmp_path, monkeypatch)
    try:
        run = PaperRun(
            paper_id="paper-verified-01",
            filename="coastal-sensor-study.pdf",
            author="Mina Shah",
            states={
                "weaknesses": StageState.COMPLETED,
                "novelty": StageState.COMPLETED,
                "clarity": StageState.PENDING,
            },
            results={
                "weaknesses": {
                    "confidence": 0.91,
                    "metrics": {"f1": 0.78},
                    "findings": [{
                        "type": "missing_baseline",
                        "description": "The analysis did not detect a baseline comparison.",
                        "evidence": {"matched_text": "We compared the proposed sensor to our own prior version."},
                        "recommended_action": "Report a comparison against a named baseline.",
                    }],
                },
                "novelty": {"findings": []},
                "clarity": {"findings": [{"type": "unclear"}]},
            },
        )
        paper_store.save(run)
        response = client.get("/api/v1/builder/papers/paper-verified-01/findings")
        assert response.status_code == 200
        findings = response.json()["findings"]
        assert len(findings) == 1
        finding = findings[0]
        assert finding["title"] == "missing_baseline"
        assert finding["description"] == "The analysis did not detect a baseline comparison."
        assert finding["evidence"]["matched_text"].startswith("We compared")
        assert finding["confidence"] == 0.91
        assert finding["recommended_action"] == "Report a comparison against a named baseline."
        assert finding["affected_section"] is None
        assert "fabricat" not in finding["description"].lower()

        dashboard = client.get("/api/v1/builder/papers/paper-verified-01/quality")
        assert dashboard.status_code == 200
        assert dashboard.json()["findings_count"] == 1
        weakness_module = next(module for module in dashboard.json()["modules"] if module["module"] == "weaknesses")
        assert weakness_module["findings_count"] == 1
        assert weakness_module["metrics"] == {"f1": 0.78}
        assert client.get("/api/v1/builder/papers/not-present/findings").status_code == 404

        draft = client.post("/api/v1/research-papers", json={
            "title": "Coastal sensor calibration",
            "paper_id": "paper-verified-01",
            "format": "generic",
        }).json()
        existing = client.post(f"/api/v1/research-papers/{draft['id']}/analyze")
        assert existing.status_code == 200
        assert existing.json()["analysis_results"]["finding_count"] == 1
        assert existing.json()["analysis_results"]["results"]["weaknesses"]["metrics"] == {"f1": 0.78}
    finally:
        client.close()
        engine.dispose()


def test_generated_paper_runs_persisted_analysis_actions_and_reanalysis(tmp_path, monkeypatch):
    client, engine, _ = make_client(tmp_path, monkeypatch)
    try:
        created = client.post("/api/v1/research-papers", json={
            "title": "A field study of coastal sensor calibration",
            "authors": "Mina Shah",
            "abstract": "This study examines calibration drift using field observations.",
            "introduction": "We propose a novel calibration method for long-term coastal sensor monitoring.",
            "related_work": "Prior work is described by the researcher.",
            "methodology": "Sensors were calibrated monthly and the observations were grouped by month.",
            "dataset": "The researcher supplied a set of field observations.",
            "experiments_results": "The supplied evaluation describes measured sensor behavior.",
            "limitations": "The observations cover a limited set of coastal stations.",
            "format": "generic",
        })
        assert created.status_code == 201, created.text
        draft_id = created.json()["paper_id"]
        generated = client.post(f"/api/v1/research-papers/{draft_id}/generate")
        assert generated.status_code == 200, generated.text

        started = client.post(f"/api/v1/research-papers/{draft_id}/analyze")
        assert started.status_code == 200, started.text
        run_id = started.json()["analysis_run_id"]
        assert started.json()["analysis_status"] in {"running", "completed"}

        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            result = client.get(f"/api/v1/research-papers/{draft_id}/analysis/{run_id}")
            assert result.status_code == 200, result.text
            if result.json()["analysis_status"] != "running":
                break
            time.sleep(0.05)
        analysis = result.json()
        assert analysis["analysis_status"] == "completed"
        assert all(analysis["modules"][name]["status"] == "completed" for name in (
            "related_work", "novelty", "weaknesses", "clarity", "reviewer_feedback"
        ))
        assert "monthly" in analysis["paper_text"]
        assert analysis["results"]["related_work"]["evidence"] == []
        assert "not available yet" in " ".join(analysis["results"]["related_work"]["limitations"]).lower()
        assert "novelty_score" not in analysis["results"]["novelty"]["metrics"]
        assert "overall_assessment" not in analysis["results"]["reviewer_feedback"]
        assert all(finding["status"] == "requires_investigation" for finding in analysis["findings"])
        assert all(finding["confidence"] is None for finding in analysis["findings"])
        novelty_finding = next(item for item in analysis["findings"] if item["module"] == "novelty")
        assert novelty_finding["evidence"]["exact_match"] is True
        assert "novel calibration method" in novelty_finding["evidence_text"].lower()
        weakness_finding = next(item for item in analysis["findings"] if item["module"] == "weaknesses")
        assert weakness_finding["recommended_action"]

        assert analysis["findings"]
        finding = analysis["findings"][0]
        added = client.post(
            f"/api/v1/research-papers/{draft_id}/analysis/{run_id}/actions",
            json={"finding_id": finding["id"]},
        )
        assert added.status_code == 201, added.text
        assert added.json()["status"] == "not_started"
        action_id = added.json()["action_id"]
        updated_action = client.patch(
            f"/api/v1/research-papers/{draft_id}/analysis/{run_id}/actions/{action_id}",
            json={"status": "in_progress"},
        )
        assert updated_action.status_code == 200
        assert updated_action.json()["status"] == "in_progress"

        edited = client.put(f"/api/v1/research-papers/{draft_id}", json={
            **created.json()["details"],
            "generated_content": {
                **generated.json()["generated_content"],
                "methodology": "Sensors were calibrated monthly against a documented reference baseline.",
            },
        })
        assert edited.status_code == 200, edited.text
        reanalysis = client.post(f"/api/v1/research-papers/{draft_id}/analyze")
        assert reanalysis.status_code == 200, reanalysis.text
        next_run_id = reanalysis.json()["analysis_run_id"]
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            refreshed = client.get(f"/api/v1/research-papers/{draft_id}/analysis/{next_run_id}")
            assert refreshed.status_code == 200, refreshed.text
            if refreshed.json()["analysis_status"] != "running":
                break
            time.sleep(0.05)
        assert refreshed.json()["analysis_status"] == "completed"
        assert refreshed.json()["version"] == edited.json()["version"]
        assert set(refreshed.json()["delta"]) == {"new", "remaining", "resolved", "changed"}
        report = client.get(
            f"/api/v1/research-papers/{draft_id}/analysis/{next_run_id}/report"
        )
        assert report.status_code == 200, report.text
        assert report.json()["summary"]["total_findings"] == refreshed.json()["finding_count"]
        report_pdf = client.get(
            f"/api/v1/research-papers/{draft_id}/analysis/{next_run_id}/report/export?format=pdf"
        )
        assert report_pdf.status_code == 200, report_pdf.text
        assert report_pdf.headers["content-type"].startswith("application/pdf")
        assert "Research Analysis Report" in "\n".join(
            page.extract_text() or "" for page in PdfReader(BytesIO(report_pdf.content)).pages
        )
        report_docx = client.get(
            f"/api/v1/research-papers/{draft_id}/analysis/{next_run_id}/report/export?format=docx"
        )
        assert report_docx.status_code == 200, report_docx.text
        docx_text = "\n".join(
            paragraph.text for paragraph in Document(BytesIO(report_docx.content)).paragraphs
        )
        assert "Research Analysis Report" in docx_text
        assert "Findings" in docx_text
        report_markdown = client.get(
            f"/api/v1/research-papers/{draft_id}/analysis/{next_run_id}/report/export?format=markdown"
        )
        assert report_markdown.status_code == 200, report_markdown.text
        assert "# Research Analysis Report:" in report_markdown.text
        if report.json()["findings"]:
            assert report.json()["findings"][0]["title"] in report_markdown.text
        else:
            assert "No findings were returned" in report_markdown.text
        assert client.get(f"/api/v1/research-papers/{draft_id}").json()["analysis_results"]["analysis_run_id"] == next_run_id

        retried = client.post(
            f"/api/v1/research-papers/{draft_id}/analysis/{next_run_id}/retry/related_work"
        )
        assert retried.status_code == 200, retried.text
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            retry_result = client.get(f"/api/v1/research-papers/{draft_id}/analysis/{next_run_id}")
            if retry_result.json()["modules"]["related_work"]["status"] == "completed":
                break
            time.sleep(0.05)
        assert retry_result.json()["analysis_status"] == "completed"
        assert retry_result.json()["results"]["related_work"]["evidence"] == []
    finally:
        client.close()
        engine.dispose()


def test_generation_marks_absent_information_and_rejects_invalid_schema(tmp_path, monkeypatch):
    client, engine, _ = make_client(tmp_path, monkeypatch)
    try:
        missing = client.post("/api/v1/builder/drafts", json={"format": "imrad"})
        assert missing.status_code == 201
        document = missing.json()["document"]
        assert {"title", "authors", "abstract", "introduction", "methods", "results", "discussion", "conclusion",
                "references and citations"}.issubset(set(document["missing_items"]))
        rendered = client.get(f"/api/v1/builder/drafts/{missing.json()['draft_id']}/export?format=markdown").text
        assert "[Missing: title]" in rendered
        assert "[Missing: authors]" in rendered
        assert "[Missing: references and citations" in rendered
        assert client.post("/api/v1/builder/drafts", json={"format": "invented"}).status_code == 422
        assert client.post("/api/v1/builder/drafts", json={"unknown": "not allowed"}).status_code == 422
        assert client.put(f"/api/v1/builder/drafts/{missing.json()['draft_id']}", json={"title": None}).status_code == 422
    finally:
        client.close()
        engine.dispose()


def test_application_registers_builder_and_database_initializes_tables(tmp_path, monkeypatch):
    from backend.app import database
    from backend.app.main import app

    engine = create_engine(f"sqlite:///{tmp_path / 'initialized.sqlite'}")
    monkeypatch.setattr(database, "engine", engine)
    database.init_db()
    tables = set(inspect(engine).get_table_names())
    assert {"builder_drafts", "builder_draft_versions", "builder_actions"}.issubset(tables)

    registered = set(app.openapi()["paths"])
    assert "/api/v1/builder/drafts" in registered
    assert "/api/v1/builder/papers/{paper_id}/findings" in registered
    assert "/api/v1/builder/papers/{paper_id}/quality" in registered
    assert "/api/v1/builder/{draft_id}/analysis/{run_id}" in registered
    assert "/api/v1/research-papers" in registered
    engine.dispose()
